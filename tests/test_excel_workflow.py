from pathlib import Path
from tempfile import TemporaryDirectory
from datetime import datetime
import unittest

from openpyxl import Workbook, load_workbook
from app.excel_service import load_excel_plan, save_results
from app.models import InvoiceResult


class ExcelWorkflowTests(unittest.TestCase):
    def make_workbook(self, folder):
        path = Path(folder) / 'synthetic-invoices.xlsx'
        workbook = Workbook()
        overdue = workbook.active
        overdue.title = 'Vencidos'
        overdue.append(['Nº FATURA', 'EMISSÃO FATURA', 'Situação fatura no Transpofrete', 'Motivo rejeite no Transpofrete'])
        overdue.append([1001, datetime(2026, 1, 10), None, 'Keep this synthetic note'])
        overdue.append([1002, datetime(2026, 1, 1), 'Already reviewed', 'Unchanged'])
        overdue.append([1003, 'not a date', None, None])
        upcoming = workbook.create_sheet('A Vencer')
        upcoming.append(['Nº FATURA', 'EMISSÃO FATURA', 'Status TSP', 'OBS.'])
        upcoming.append([1001, datetime(2026, 1, 5), None, None])
        workbook.save(path)
        return path

    def test_duplicate_invoices_share_a_task_and_completed_rows_are_skipped(self):
        with TemporaryDirectory() as folder:
            plan = load_excel_plan(str(self.make_workbook(folder)))
            self.assertEqual(len(plan.tasks), 1)
            self.assertEqual(plan.tasks[0].invoice, '1001')
            self.assertEqual(len(plan.tasks[0].targets), 2)
            self.assertEqual(plan.dt_from, datetime(2026, 1, 5))

    def test_results_update_all_targets_without_touching_completed_rows(self):
        with TemporaryDirectory() as folder:
            path = self.make_workbook(folder)
            plan = load_excel_plan(str(path))
            save_results(plan, {'1001': InvoiceResult(status='Demo reviewed', motivo='Synthetic reason')})
            workbook = load_workbook(path)
            self.assertEqual(workbook['Vencidos']['C2'].value, 'Demo reviewed')
            self.assertEqual(workbook['A Vencer']['C2'].value, 'Demo reviewed')
            self.assertEqual(workbook['Vencidos']['D2'].value, 'Synthetic reason')
            self.assertEqual(workbook['Vencidos']['C3'].value, 'Already reviewed')

    def test_empty_result_reason_preserves_existing_notes(self):
        with TemporaryDirectory() as folder:
            path = self.make_workbook(folder)
            plan = load_excel_plan(str(path))
            save_results(plan, {'1001': InvoiceResult(status='Demo reviewed')})
            self.assertEqual(load_workbook(path)['Vencidos']['D2'].value, 'Keep this synthetic note')

    def test_missing_required_headers_fail_without_modifying_input(self):
        with TemporaryDirectory() as folder:
            path = Path(folder) / 'invalid.xlsx'
            workbook = Workbook()
            workbook.active.title = 'Vencidos'
            workbook.active.append(['Unexpected column'])
            workbook.save(path)
            before = path.read_bytes()
            with self.assertRaises(ValueError):
                load_excel_plan(str(path))
            self.assertEqual(path.read_bytes(), before)

if __name__ == '__main__':
    unittest.main()
