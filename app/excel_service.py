from __future__ import annotations

import re
import unicodedata
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .config import EXCEL, SheetExcelConfig
from .models import ExcelPlan, InvoiceResult, InvoiceTarget, InvoiceTask


def _normalize_text(value: Any) -> str:
    text = str(value or "").replace("\xa0", " ").strip()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.lower()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _is_blank(value: Any) -> bool:
    if value is None:
        return True

    text = str(value).strip()
    return text == ""


def _invoice_to_text(value: Any) -> str:
    if value is None:
        return ""

    if isinstance(value, int):
        return str(value).strip()

    if isinstance(value, float):
        if value.is_integer():
            return str(int(value)).strip()
        return str(value).strip()

    return str(value).strip()


def _parse_date(value: Any) -> datetime | None:
    if value is None:
        return None

    if isinstance(value, datetime):
        return value

    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)

    text = str(value).strip()

    if not text:
        return None

    formats = (
        "%d/%m/%Y %H:%M:%S",
        "%d/%m/%Y %H:%M",
        "%d/%m/%Y",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d",
    )

    for fmt in formats:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass

    return None


def _find_col(headers: dict[int, Any], aliases: tuple[str, ...], sheet_name: str) -> int:
    normalized_headers = {
        _normalize_text(header): col
        for col, header in headers.items()
        if not _is_blank(header)
    }

    for alias in aliases:
        alias_norm = _normalize_text(alias)
        if alias_norm in normalized_headers:
            return normalized_headers[alias_norm]

    for alias in aliases:
        alias_norm = _normalize_text(alias)

        for header_norm, col in normalized_headers.items():
            if alias_norm in header_norm or header_norm in alias_norm:
                return col

    raise ValueError(
        f'Na sheet "{sheet_name}", não encontrei uma coluna obrigatória. '
        f"Aliases esperados: {', '.join(aliases)}"
    )


def _read_sheet_targets(
    wb,
    sheet_config: SheetExcelConfig,
    grouped_targets: dict[str, list[InvoiceTarget]],
    dates: list[datetime],
) -> int:
    if sheet_config.sheet_name not in wb.sheetnames:
        return 0

    ws = wb[sheet_config.sheet_name]

    headers = {
        col_idx: ws.cell(row=1, column=col_idx).value
        for col_idx in range(1, ws.max_column + 1)
    }

    invoice_col = _find_col(
        headers=headers,
        aliases=sheet_config.col_invoice_aliases,
        sheet_name=sheet_config.sheet_name,
    )

    emissao_col = _find_col(
        headers=headers,
        aliases=sheet_config.col_emissao_fatura_aliases,
        sheet_name=sheet_config.sheet_name,
    )

    status_col = _find_col(
        headers=headers,
        aliases=sheet_config.col_status_aliases,
        sheet_name=sheet_config.sheet_name,
    )

    motivo_col = _find_col(
        headers=headers,
        aliases=sheet_config.col_motivo_aliases,
        sheet_name=sheet_config.sheet_name,
    )

    pending_rows = 0

    for row in range(2, ws.max_row + 1):
        invoice = _invoice_to_text(ws.cell(row=row, column=invoice_col).value)

        if not invoice or invoice == "-":
            continue

        current_status = ws.cell(row=row, column=status_col).value

        # Mantém a regra anterior:
        # só consulta linhas ainda sem status preenchido.
        if not _is_blank(current_status):
            continue

        emissao = _parse_date(ws.cell(row=row, column=emissao_col).value)

        if emissao is None:
            continue

        grouped_targets[invoice].append(
            InvoiceTarget(
                sheet_name=sheet_config.sheet_name,
                row=row,
                status_col=status_col,
                motivo_col=motivo_col,
            )
        )

        dates.append(emissao)
        pending_rows += 1

    return pending_rows


def load_excel_plan(file_path: str) -> ExcelPlan:
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"Arquivo não encontrado: {file_path}")

    keep_vba = path.suffix.lower() == ".xlsm"
    wb = load_workbook(file_path, keep_vba=keep_vba)

    grouped_targets: dict[str, list[InvoiceTarget]] = defaultdict(list)
    dates: list[datetime] = []
    processed_sheets: list[str] = []
    total_pending_rows = 0

    for sheet_config in EXCEL.sheets:
        if sheet_config.sheet_name not in wb.sheetnames:
            continue

        pending_rows = _read_sheet_targets(
            wb=wb,
            sheet_config=sheet_config,
            grouped_targets=grouped_targets,
            dates=dates,
        )

        processed_sheets.append(sheet_config.sheet_name)
        total_pending_rows += pending_rows

    if not processed_sheets:
        expected = ", ".join(sheet.sheet_name for sheet in EXCEL.sheets)
        raise ValueError(
            "Nenhuma das sheets esperadas foi encontrada. "
            f"Sheets esperadas: {expected}"
        )

    if not grouped_targets:
        raise ValueError(
            "Não encontrei faturas pendentes nas sheets configuradas. "
            "Verifique se as colunas de status estão vazias nas linhas desejadas."
        )

    if not dates:
        raise ValueError(
            'Não encontrei datas válidas na coluna "EMISSÃO FATURA" '
            "para as faturas pendentes."
        )

    tasks = [
        InvoiceTask(invoice=invoice, targets=targets)
        for invoice, targets in grouped_targets.items()
    ]

    return ExcelPlan(
        file_path=file_path,
        sheet_names=processed_sheets,
        dt_from=min(dates),
        tasks=tasks,
    )


def save_results(plan: ExcelPlan, results: dict[str, InvoiceResult]) -> None:
    path = Path(plan.file_path)
    keep_vba = path.suffix.lower() == ".xlsm"

    wb = load_workbook(plan.file_path, keep_vba=keep_vba)

    for task in plan.tasks:
        result = results.get(
            task.invoice,
            InvoiceResult(status=EXCEL.value_not_found),
        )

        for target in task.targets:
            if target.sheet_name not in wb.sheetnames:
                continue

            ws = wb[target.sheet_name]

            ws.cell(row=target.row, column=target.status_col).value = result.status

            # Só preenche observação/motivo quando houver motivo capturado.
            # Assim evita apagar OBS. ou Motivo já existente.
            if result.motivo:
                ws.cell(row=target.row, column=target.motivo_col).value = result.motivo

    wb.save(plan.file_path)